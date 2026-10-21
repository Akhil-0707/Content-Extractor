"""Parsers for Word documents and PowerPoint decks."""

from __future__ import annotations

from app.ingest.base import Block, BlockKind, ParsedDocument, ParseError

# Word's built-in heading styles are named "Heading 1" … "Heading 9"
HEADING_STYLE_PREFIX = "heading"


def parse_docx(raw: bytes) -> ParsedDocument:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - installed in the image
        raise ParseError("DOCX support requires python-docx") from exc

    import io

    try:
        document = docx.Document(io.BytesIO(raw))
    except Exception as exc:
        raise ParseError(f"Not a readable Word document: {exc}") from exc

    blocks: list[Block] = []
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name if paragraph.style else "").lower()
        kind = BlockKind.HEADING if style.startswith(HEADING_STYLE_PREFIX) else BlockKind.TEXT
        blocks.append(Block(kind, text, {"paragraph": index}))

    for table_index, table in enumerate(document.tables):
        rendered = _render_table([[cell.text for cell in row.cells] for row in table.rows])
        if rendered:
            blocks.append(Block(BlockKind.TABLE, rendered, {"table": table_index}))

    title = (document.core_properties.title or "").strip() or None
    if not title:
        title = next((b.text for b in blocks if b.kind is BlockKind.HEADING), None)

    return ParsedDocument(
        blocks=blocks,
        title=title,
        meta={"source": "docx", "paragraphs": len(document.paragraphs), "tables": len(document.tables)},
    )


def parse_pptx(raw: bytes) -> ParsedDocument:
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover - installed in the image
        raise ParseError("PPTX support requires python-pptx") from exc

    import io

    try:
        presentation = Presentation(io.BytesIO(raw))
    except Exception as exc:
        raise ParseError(f"Not a readable PowerPoint file: {exc}") from exc

    blocks: list[Block] = []
    for number, slide in enumerate(presentation.slides, start=1):
        ref = {"slide": number}
        title_shape = slide.shapes.title
        if title_shape is not None and (title_shape.text or "").strip():
            blocks.append(Block(BlockKind.HEADING, title_shape.text, dict(ref)))

        for shape in slide.shapes:
            if shape is title_shape or not shape.has_text_frame:
                continue
            text = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
            if text.strip():
                blocks.append(Block(BlockKind.TEXT, text, dict(ref)))

        # Speaker notes often hold the trainer's actual explanation
        if slide.has_notes_slide:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip()
            if notes:
                blocks.append(Block(BlockKind.TEXT, notes, {**ref, "notes": True}))

    title = next((b.text for b in blocks if b.kind is BlockKind.HEADING), None)
    return ParsedDocument(
        blocks=blocks,
        title=title,
        meta={"source": "pptx", "slides": len(presentation.slides)},
    )


def _render_table(rows: list[list[str]]) -> str:
    """Render a table as pipe-separated rows, which models read reliably."""
    cleaned = [[" ".join(cell.split()) for cell in row] for row in rows]
    cleaned = [row for row in cleaned if any(row)]
    return "\n".join(" | ".join(row) for row in cleaned)
